#!/usr/bin/env python3
"""Deterministic checks for 7Sec outbound documents (contracts, appendices/order forms, annexes).

Usage: python3 check_doc.py <file.pdf|file.docx> [--json]

Prints findings grouped by severity. Exit code 1 if any BLOCK finding.
Rules are documented in ../references/rules.md; keep both in sync.
"""
import json
import re
import shutil
import subprocess
import sys
from datetime import date

SEVEN_REG = "561824m"
SEVEN_VAT = "ATU77298229"
SEVEN_ADDR = "Hiessgasse 12/3"

findings = []


def add(sev, rule, msg, where=""):
    findings.append({"severity": sev, "rule": rule, "message": msg, "where": where})


def extract_text(path):
    low = path.lower()
    if low.endswith(".pdf"):
        if shutil.which("pdftotext"):
            return subprocess.run(["pdftotext", "-layout", path, "-"], capture_output=True, text=True, check=True).stdout
        from pypdf import PdfReader  # fallback
        return "\n".join(p.extract_text() or "" for p in PdfReader(path).pages)
    if low.endswith((".docx", ".dotx")):
        if shutil.which("pandoc"):
            body = subprocess.run(["pandoc", "-t", "plain", "--wrap=none", path], capture_output=True, text=True, check=True).stdout
        else:
            import docx  # python-docx fallback
            d = docx.Document(path)
            parts = [p.text for p in d.paragraphs]
            for t in d.tables:
                for r in t.rows:
                    parts.append(" | ".join(c.text for c in r.cells))
            body = "\n".join(parts)
        return body + "\n--- headers and footers ---\n" + docx_headers_footers(path)
    raise SystemExit("Unsupported file type: use .pdf or .docx")


def docx_headers_footers(path):
    import zipfile
    out = []
    with zipfile.ZipFile(path) as z:
        for n in sorted(z.namelist()):
            if re.fullmatch(r"word/(header|footer)\d*\.xml", n):
                x = z.read(n).decode("utf8", "ignore")
                x = re.sub(r"</w:p>", "\n", x)
                x = re.sub(r"<[^>]+>", "", x)
                out.append(x)
    return "\n".join(out)


def parse_date(s):
    d, m, y = s.split(".")
    return date(int(y), int(m), int(d))


def check_placeholders(flat, text):
    for m in re.finditer(r"_{3,}", text):
        ctx = re.sub(r"\s+", " ", text[max(0, m.start() - 40): m.end() + 20])
        add("BLOCK", "placeholder", "Unfilled blank (underscores) left in the document.", ctx)
        break
    for pat, label in [(r"\[[^\]\n]{1,60}\]", "text in [brackets]"), (r"\bXXX+\b", "XXX"), (r"\bTBD\b", "TBD"),
                       (r"YYYYMMDD|DD\.MM\.YYYY|DAY/MONTH/YEAR", "date placeholder"), (r"(?i)lorem ipsum", "lorem ipsum"),
                       (r"CLIENT[’']S NAME|to Client[’']s name \(“Client”\)", "CLIENT'S NAME"), (r"(?i)\d+ sentences about the client", "'sentences about the Client' prompt"),
                       (r"(?i)\(alter as needed\)", "'(alter as needed)'"), (r"…|\.{4,}", "ellipsis (…) left to fill in"),
                       (r"(?i)Discount\s*\(if applicable\)", "'Discount (if applicable)'"), (r"TRADE NAME", "TRADE NAME")]:
        m = re.search(pat, flat)
        if m:
            add("BLOCK", "placeholder", f"Possible placeholder left in: {label}.", flat[max(0, m.start() - 40): m.end() + 20])


def check_numbers(flat, has_contract_body):
    tok = re.compile(r"\b([A-Z]{2,}[A-Z0-9]*\d+[A-Z]{0,4}/\d+(?:/\d+)*)\b")
    fmt = re.compile(r"^[A-Z][A-Z0-9]*?(\d{6})(\d{2})([A-Z]{0,4})/(\d{6})(\d{2})$")
    contract_nos, order_nos = set(), set()
    for m in tok.finditer(flat):
        t = m.group(1)
        before = flat[max(0, m.start() - 14): m.start()]
        if re.search(r"Order\s*No\.?:?\s*$", before, re.I):
            order_nos.add(t)
        else:
            contract_nos.add(t)
    # contract references must have exactly two parts
    clean_contracts = set()
    for c in contract_nos:
        parts = c.split("/")
        if len(parts) > 2:
            base = "/".join(parts[:2])
            add("BLOCK", "contract-number", f"Contract number '{c}' has an order suffix '/{'/'.join(parts[2:])}'. "
                f"The contract number is two parts ('{base}'); the /N suffix belongs only to the Order No.", c)
            c = base
        clean_contracts.add(c)
    if len(clean_contracts) > 1:
        add("BLOCK", "contract-number", f"More than one contract number appears: {sorted(clean_contracts)}.")
    for c in clean_contracts:
        m = fmt.match(c)
        if not m:
            add("BLOCK", "contract-number", f"Contract number '{c}' does not match CLIENT+YYMMDD+NN[suffix]/YYMMDDNN.", c)
            continue
        yy, mm, dd = m.group(4)[:2], m.group(4)[2:4], m.group(4)[4:6]
        try:
            num_date = date(2000 + int(yy), int(mm), int(dd))
        except ValueError:
            add("BLOCK", "contract-number", f"Second part of '{c}' is not a valid date (YYMMDD).", c)
            continue
        # date stated in the document
        stated = re.findall(r"\bdated\s+(\d{2}\.\d{2}\.\d{4})", flat)
        if not stated and has_contract_body:
            stated = re.findall(r"On\s+this\s+date\s+(\d{2}\.\d{2}\.\d{4})", flat)
        for s in stated[:1]:
            if parse_date(s) != num_date:
                add("FIX", "contract-date", f"Contract number {c} implies date {num_date:%d.%m.%Y}, but the document says {s}.", c)
    for o in order_nos:
        parts = o.split("/")
        if len(parts) != 3:
            add("BLOCK", "order-number", f"Order No. '{o}' should be contract number + '/N'.", o)
            continue
        base = "/".join(parts[:2])
        if clean_contracts and base not in clean_contracts:
            add("BLOCK", "order-number", f"Order No. '{o}' does not start with the contract number {sorted(clean_contracts)}.", o)
    if re.search(r"Order\s*No\.?:?\s*[A-Z_]+[A-Z0-9_]*/[_\d]+\d\b(?!/)", flat) and not order_nos:
        pass
    # proposal number vs first part
    pm = re.search(r"Proposal\s+No\.?:?\s*([A-Z][A-Z0-9]+\d[A-Z]{0,4})\b", flat)
    if pm and clean_contracts:
        first = {c.split("/")[0] for c in clean_contracts}
        if pm.group(1) not in first:
            add("FIX", "proposal-number", f"Proposal No. {pm.group(1)} differs from the first part of the contract number {sorted(first)}.")
    return clean_contracts, order_nos


def check_parties(flat):
    if "7Security GmbH" not in flat:
        return
    for m in re.finditer(r"(?:registration\s+(?:number|No\.?)\s*)((?:FN\s*)?5618\d\d\s?m)", flat):
        if m.group(1) != SEVEN_REG:
            add("BLOCK", "7sec-details", f"7Security registration number written '{m.group(1)}'; standard is '{SEVEN_REG}'.")
    for m in re.finditer(r"7Security GmbH[^,]{0,50},\s*with registration (?:number|No\.?)\s*([^,]+?),", flat):
        if m.group(1).strip() != SEVEN_REG:
            add("BLOCK", "7sec-details", f"7Security registration number is written '{m.group(1).strip()}'; it should be '{SEVEN_REG}'.", m.group(0)[:90])
    if "ATU" in flat and SEVEN_VAT not in flat:
        add("BLOCK", "7sec-details", f"7Security VAT number differs from {SEVEN_VAT}.")
    if re.search(r"integral part|is entered between", flat):
        if SEVEN_VAT not in flat:
            add("BLOCK", "7sec-details", f"7Security VAT number {SEVEN_VAT} not found in the preamble.")
        if SEVEN_ADDR not in flat:
            add("FIX", "7sec-details", f"7Security legal address ({SEVEN_ADDR}) is missing from the preamble; the templates include it.")
    # client name in signature block must also be the client named in the preamble
    for m in re.finditer(r"\bFor\s+([A-Z][^:\n]{2,60}?):", flat):
        name = m.group(1).strip()
        if "7Security" in name:
            continue
        if flat.lower().count(name.lower()) < 2:
            add("BLOCK", "signatory-party", f"Signature block is for '{name}', which is not named elsewhere in the document "
                "(the preamble names a different party?).", name)


def check_bulgarian(flat):
    sub = "7Security Ltd" in flat
    bg_law = bool(re.search(r"laws? of the Republic of Bulgaria", flat))
    at_law = bool(re.search(r"laws? of the Republic of Austria", flat))
    if bg_law and at_law:
        add("FIX", "governing-law", "Both Bulgarian and Austrian law are named.")
    if sub and at_law and not bg_law:
        add("FIX", "bulgarian-version", "7Security Ltd (Bulgarian subcontractor) appears, but governing law is Austria. "
            "Bulgarian-client templates use Bulgarian law and courts.")
    if bg_law and not sub:
        add("FIX", "bulgarian-version", "Governing law is Bulgaria but 7Security Ltd (subcontractor) is not named.")


def check_totals(text):
    """Check each order-form block: items minus discounts must equal its TOTAL line.
    A final 'TOTAL CONTRACT COST' must equal the sum of the yearly totals."""
    if not re.search(r"Order Form|COMMERCIAL PROPOSAL|TOTAL COST", text):
        return
    items, discounts, totals = 0, 0, []
    seen_item = False
    for line in text.splitlines():
        ams = re.findall(r"(-\s*)?€\s*([\d,]+)", line)
        if not ams:
            continue
        val = int(ams[-1][1].replace(",", ""))
        if re.search(r"TOTAL|ORDER COST|COST PER MONTH", line, re.I):
            if re.search(r"CONTRACT", line, re.I) and totals:
                if sum(t for t, _ in totals) != val:
                    add("BLOCK", "totals", f"Yearly/order totals add up to €{sum(t for t, _ in totals):,}, but the contract total says €{val:,}.", line.strip())
                items = discounts = 0
                continue
            expect = items - discounts
            if seen_item and expect != val:
                add("BLOCK", "totals", f"Line items minus discounts = €{expect:,}, but the total says €{val:,}.", line.strip())
            totals.append((val, line.strip()))
            items = discounts = 0
            seen_item = False
        elif re.search(r"Discount", line, re.I):
            discounts += val
        else:
            items += val
            seen_item = True


WORDS = {"one": 1, "two": 2, "three": 3, "four": 4, "five": 5, "six": 6, "seven": 7, "eight": 8, "nine": 9, "ten": 10,
         "twelve": 12, "fourteen": 14, "fifteen": 15, "twenty": 20, "thirty": 30, "forty": 40, "sixty": 60, "ninety": 90,
         "twenty-four": 24, "twenty four": 24, "thirty-six": 36}


def check_number_words(flat):
    for m in re.finditer(r"\b(\d+)\s*\(\s*([A-Za-z -]+?)\s*\)", flat):
        w = m.group(2).lower().strip()
        if w in WORDS and WORDS[w] != int(m.group(1)):
            add("BLOCK", "number-words", f"'{m.group(0)}': the number and the word in brackets disagree.", m.group(0))


def check_validity(flat):
    v = re.search(r"offer stands for\s+(\d+)", flat, re.I)
    d = re.search(r"ISSUED:?\s*(\d{1,2})[./](\d{1,2})[./](\d{4})", flat)
    if v and d:
        issued = date(int(d.group(3)), int(d.group(2)), int(d.group(1)))
        left = (issued.toordinal() + int(v.group(1))) - date.today().toordinal()
        if left < 0:
            add("BLOCK", "offer-validity", f"The offer was issued {issued:%d.%m.%Y} and stands for {v.group(1)} days, so it expired {-left} days ago.")
        elif left < 7:
            add("INFO", "offer-validity", f"The offer expires in {left} days.")


SERVICE_WORDS = {
    "SWIFT": {"ISMS", "AICPA", "Trust Services", "PSD2"},
    "PSD2": {"ISMS", "AICPA", "Trust Services", "SWIFT"},
    "SOC2": {"ISMS", "International Organization for Standardization", "SWIFT", "PSD2"},
    "ISO": {"AICPA", "Trust Services", "SWIFT", "PSD2"},
}


def check_service_leftovers(flat):
    """Flag wording that belongs to another service's template (e.g. ISMS in a SWIFT offer).
    A combined offer (e.g. ISO + SOC2) only gets terms that are foreign to every service in it."""
    head = flat[:700]
    declared = [n for pat, n in [(r"SWIFT", "SWIFT"), (r"PSD2", "PSD2"), (r"SOC ?2", "SOC2"), (r"ISO/IEC 27001|ISO 27001", "ISO")]
                if re.search(pat, head, re.I)]
    if not declared:
        return
    # the standard 'About Us' and team-credentials sentences list every service; ignore them
    body = re.sub(r"We provide professional consulting services in compliance auditing, such as.*?(?<=[a-z0-9])\. ", " ", flat)
    body = re.sub(r"[^.]{0,400}\bOSCP\b[^.]{0,80}\.", " ", body)
    forbidden = set.intersection(*(SERVICE_WORDS[d] for d in declared))
    label = " + ".join(declared)
    for term in sorted(forbidden):
        m = re.search(r"\b" + re.escape(term) + r"\b", body)
        if m:
            add("FIX", "wrong-service-text", f"This is a {label} document but it mentions '{term}'. Copied from another service's template?",
                body[max(0, m.start() - 50): m.end() + 40])
    if re.search(r"QSA company", body):
        add("FIX", "wrong-service-text", f"'QSA company' (PCI wording) appears in a {label} document.", "QSA company")


def check_missing_prices(text):
    if re.search(r"€\s*(?:\||$)", text, re.M):
        add("BLOCK", "prices", "A price field is empty (a '€' with no amount).")


def check_offer_leftovers(flat):
    m = re.search(r"€\s*\d+/\d[\d,]*", flat)
    if m:
        add("BLOCK", "prices", f"Unresolved either/or amount '{m.group(0)}'. Pick one.", m.group(0))
    m = re.search(r"\(\s*NOTE:[^)]{0,160}\)?", flat)
    if m:
        add("BLOCK", "template-instruction", "An internal instruction from the template is still in the text.", m.group(0)[:120])
    m = re.search(r"\(OPTIONAL\)", flat)
    if m:
        add("FIX", "template-instruction", "'(OPTIONAL)' is still in a heading: keep or remove the section, and take the word out.")
    tags = {t for t in ("L1", "ML2", "SPL2") if re.search(r"\b" + t + r"\b[^€]{0,12}:\s*€", flat)}
    if len(tags) > 1:
        add("FIX", "price-options", f"Several reporting price options are left in ({', '.join(sorted(tags))}). Keep only the one that applies.")
    years = [int(y) for y in re.findall(r"\b(?:(?:1st|2nd|3rd)\s+)?Year\s+(20\d\d)", flat)]
    if years:
        d = re.search(r"ISSUED:?\s*(\d{1,2})[./](\d{1,2})[./](\d{4})", flat)
        ref = int(d.group(3)) if d else date.today().year
        if years[0] < ref:
            add("FIX", "year-labels", f"The first contract year is labelled {years[0]}, but the offer is from {ref}. Update the year labels.")
        if any(b - a != 1 for a, b in zip(years, years[1:])):
            add("FIX", "year-labels", f"The year labels are not consecutive: {years}.")


def check_prepayment_wording(flat):
    """Decision: prepayment/advance invoices are due from the ISSUANCE date, in offers and contracts alike."""
    m = re.search(r"advance invoice within 10 \(ten\) calendar days from its receipt|prepayment invoice[^.]{0,60}from (?:its|the) receipt", flat, re.I)
    if m:
        add("FIX", "payment-terms", "The advance/prepayment invoice is due 'from its receipt'. The rule is 10 calendar days from the "
            "ISSUANCE date, as in the offers.", m.group(0)[:100])


def check_units(flat):
    m = re.search(r"~?\b1\s+(Days|Months|Weeks|Years)\b", flat)
    if m:
        add("INFO", "typo", f"Number and unit disagree: '{m.group(0)}'.", m.group(0))
    m = re.search(r"~?\b([2-9]|\d{2,})\s+(Day|Month|Week)\b(?!s)", flat)
    if m:
        add("INFO", "typo", f"Number and unit disagree: '{m.group(0)}'.", m.group(0))


def check_offer_code(flat):
    """Offer cover: CODE = CLIENT + YYMMDD (issue date) + NN + service suffix, e.g. TENEN26101310SWIFT."""
    cover = flat[:900]
    if "--- headers and footers ---" in flat:  # Word files keep the cover fields in the footer
        cover += " " + flat.split("--- headers and footers ---")[-1]
    d = re.search(r"ISSUED:?\s*(?:VERSION:?\s*)?(\d{1,2})[./](\d{1,2})[./](\d{4})", cover)
    def real_date(ymd):
        try:
            date(2000 + int(ymd[:2]), int(ymd[2:4]), int(ymd[4:6]))
            return True
        except ValueError:
            return False
    toks = [t for t in re.findall(r"\b([A-Z][A-Z0-9]*?)(\d{6})(\d{2})([A-Z]{0,8})\b", cover)
            if real_date(t[1]) and t[0] not in ("ATU", "BG", "CY", "LU", "DE")]
    if not toks:
        if re.search(r"\bCODE:", cover):
            add("BLOCK", "offer-code", "The offer cover has no code (expected CLIENT + issue date YYMMDD + number + service, e.g. TENEN26101310SWIFT).")
        return
    name, ymd, seq, suffix = toks[0]
    if d:
        issued = f"{d.group(3)[2:]}{int(d.group(2)):02d}{int(d.group(1)):02d}"
        if issued != ymd:
            add("BLOCK", "offer-code", f"The code {name}{ymd}{seq}{suffix} contains date {ymd}, but the offer was issued on {d.group(1)}.{d.group(2)}.{d.group(3)}.")


def check_typos(flat):
    for m in re.finditer(r"(?<![.\d])\.\.(?!\.)", flat):
        add("INFO", "typo", "Double full stop.", flat[max(0, m.start() - 30): m.end() + 10])
        break
    for m in re.finditer(r"\b([A-Za-z]{3,})\s+\1\b", flat):
        add("INFO", "typo", f"Repeated word '{m.group(1)}'.", flat[max(0, m.start() - 25): m.end() + 15])
        break


def main():
    args = [a for a in sys.argv[1:] if not a.startswith("--")]
    if not args:
        raise SystemExit(__doc__)
    path = args[0]
    text = extract_text(path)
    flat = re.sub(r"\s+", " ", text)
    has_contract_body = bool(re.search(r"Confidentiality and Data Protection", flat))
    kinds = []
    if has_contract_body:
        kinds.append("contract")
    if re.search(r"\bAPPENDIX\s+[IVX\d]+", flat):
        kinds.append("appendix/order form")
    if re.search(r"\bAnnex\s+[IVX\d]+", flat):
        kinds.append("annex")
    is_offer = bool(re.search(r"COMMERCIAL PROPOSAL|\(“proposal”, “offer”", flat))
    if is_offer:
        kinds.append("offer")
    check_placeholders(flat, text)
    check_numbers(flat, has_contract_body)
    check_number_words(flat)
    check_prepayment_wording(flat)
    check_validity(flat)
    check_service_leftovers(flat)
    if is_offer:
        check_missing_prices(text)
        check_offer_leftovers(flat)
        check_offer_code(flat)
        check_units(flat)
    check_parties(flat)
    check_bulgarian(flat)
    check_totals(text)
    check_typos(flat)
    order = {"BLOCK": 0, "FIX": 1, "INFO": 2}
    findings.sort(key=lambda f: order[f["severity"]])
    if "--json" in sys.argv:
        print(json.dumps({"file": path, "detected": kinds, "findings": findings}, indent=2, ensure_ascii=False))
    else:
        print(f"File: {path}\nDetected: {', '.join(kinds) or 'unknown'}")
        if not findings:
            print("No issues found by the automatic checks.")
        for f in findings:
            w = f"  [{f['where']}]" if f["where"] else ""
            print(f"{f['severity']:5} {f['rule']}: {f['message']}{w}")
    sys.exit(1 if any(f["severity"] == "BLOCK" for f in findings) else 0)


if __name__ == "__main__":
    main()
